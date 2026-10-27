import io
import re

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_TAB_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt
from fpdf import FPDF
from fpdf.enums import XPos, YPos

import redact

BULLET = re.compile(r'^\s*[-*•·▪◦]\s+')

# the core pdf fonts only cover latin-1
PLAIN = str.maketrans({'‘': "'", '’': "'", '“': '"', '”': '"', '–': '-', '—': '-', '…': '...', ' ': ' '})


DASH_ONLY = re.compile(r'^[-–—]$')
HEADING_TAIL = re.compile(r'^(.*\S)\s{2,}([A-Z][A-Z &]{3,})$')
DATES_TAIL = re.compile(r'^(.+?)\s+\|\s+([^|]*(?:(?:19|20)\d{2}|present|current)[^|]*)$', re.IGNORECASE)


def tidy(text):
    # pdf extraction glues side-by-side text onto one line and strands dashes from date ranges
    lines = []
    for line in text.splitlines():
        line = line.rstrip()
        match = HEADING_TAIL.match(line)
        if match and redact.SECTION.match(match.group(2)):
            lines += [match.group(1), match.group(2)]
        else:
            lines.append(line if line.strip() else '')

    out = []
    for line in lines:
        if out and DASH_ONLY.match(out[-1].strip()):
            out.pop()
            if out and out[-1].strip() and line.strip():
                out[-1] = f'{out[-1]} - {line.strip()}'
                continue
        out.append(line)
    if out and DASH_ONLY.match(out[-1].strip()):
        out.pop()
    return '\n'.join(out)


def is_heading(line):
    words = line.split()
    if not words or len(words) > 5:
        return False
    if redact.SECTION.match(line) or redact.REFERENCES.match(line):
        return True
    return line.isupper() and len(line) > 3


def parse(text):
    header, body = redact.split_header(tidy(text).strip())
    lines = [line.strip() for line in header.splitlines() if line.strip()]
    blocks = []
    if lines:
        blocks.append(('name', lines[0]))
        if lines[1:]:
            blocks.append(('contact', '  |  '.join(lines[1:])))

    lines = [line.strip() for line in body.splitlines()]
    for i, line in enumerate(lines):
        after = lines[i + 1] if i + 1 < len(lines) else ''
        if not line:
            if blocks and blocks[-1][0] not in ('heading', 'gap'):
                blocks.append(('gap', ''))
        elif is_heading(line):
            blocks.append(('heading', line.rstrip(':').upper()))
        elif BULLET.match(line):
            blocks.append(('bullet', BULLET.sub('', line)))
        elif blocks and blocks[-1][0] == 'bullet' and not DATES_TAIL.match(line) and not DATES_TAIL.match(after):
            # a bullet that wrapped onto the next line
            blocks[-1] = ('bullet', f'{blocks[-1][1]} {line}')
        else:
            blocks.append(('text', line))
    return entries(blocks)


def entries(blocks):
    # "Org, Place | Dates" under a title line makes a job or school entry: bold title with dates on the right
    out = []
    for kind, line in blocks:
        match = DATES_TAIL.match(line) if kind == 'text' else None
        if not match:
            out.append((kind, line))
            continue
        left, dates = match.group(1), match.group(2)
        # the title can run over a couple of lines, like a degree and its major
        run = 0
        while run < len(out) and out[-1 - run][0] == 'text':
            run += 1
        take = run if run <= 3 else 1
        title = [out.pop()[1] for _ in range(take)][::-1]
        if title:
            out.append(('role', ', '.join(title), dates))
            out.append(('org', left))
        else:
            out.append(('role', left, dates))
    return [block if len(block) == 3 else (*block, '') for block in out]


def filename(text, ext):
    blocks = parse(text)
    name = blocks[0][1] if blocks and blocks[0][0] == 'name' else ''
    name = re.sub(r'[^A-Za-z0-9]+', '_', name).strip('_')
    return f'{name}_Resume.{ext}' if name else f'Resume.{ext}'


def underline(paragraph):
    # docx has no api for paragraph borders, so set it on the xml
    border = OxmlElement('w:pBdr')
    bottom = OxmlElement('w:bottom')
    for key, value in [('val', 'single'), ('sz', '6'), ('space', '1'), ('color', '808080')]:
        bottom.set(qn(f'w:{key}'), value)
    border.append(bottom)
    paragraph._p.get_or_add_pPr().append(border)


def to_docx(text):
    doc = Document()
    for section in doc.sections:
        section.top_margin = section.bottom_margin = Inches(0.7)
        section.left_margin = section.right_margin = Inches(0.8)
    style = doc.styles['Normal']
    style.font.name = 'Calibri'
    style.font.size = Pt(11)
    style.paragraph_format.space_after = Pt(2)

    previous = None
    for kind, line, dates in parse(text):
        if kind == 'gap':
            continue
        if kind == 'bullet':
            doc.add_paragraph(line, style='List Bullet')
            previous = kind
            continue
        p = doc.add_paragraph()
        run = p.add_run(line)
        if kind == 'role':
            run.bold = True
            if dates:
                p.paragraph_format.tab_stops.add_tab_stop(Inches(6.9), WD_TAB_ALIGNMENT.RIGHT)
                p.add_run('\t' + dates)
            if previous != 'heading':
                p.paragraph_format.space_before = Pt(8)
            p.paragraph_format.space_after = Pt(0)
            p.paragraph_format.keep_with_next = True
        elif kind == 'org':
            run.italic = True
            p.paragraph_format.keep_with_next = True
        elif kind == 'name':
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            run.bold = True
            run.font.size = Pt(20)
        elif kind == 'contact':
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            run.font.size = Pt(10)
        elif kind == 'heading':
            run.bold = True
            p.paragraph_format.space_before = Pt(12)
            p.paragraph_format.space_after = Pt(4)
            p.paragraph_format.keep_with_next = True
            underline(p)
        previous = kind

    out = io.BytesIO()
    doc.save(out)
    return out.getvalue()


def to_pdf(text):
    pdf = FPDF(format='letter')
    pdf.set_margins(19, 15, 19)
    pdf.set_auto_page_break(True, margin=15)
    pdf.add_page()
    width = pdf.epw

    previous = None
    for kind, line, dates in parse(text):
        line, dates = (s.translate(PLAIN).encode('latin-1', 'replace').decode('latin-1') for s in (line, dates))
        if kind == 'role':
            if previous not in ('heading', 'gap'):
                pdf.ln(2.5)
            # dates go on the right first, then the title wraps in whatever room is left
            if pdf.get_y() > pdf.page_break_trigger - 20:
                pdf.add_page()
            y = pdf.get_y()
            pdf.set_font('Helvetica', '', 10.5)
            room = pdf.get_string_width(dates) + 4 if dates else 0
            if dates:
                pdf.set_xy(pdf.l_margin + width - room, y)
                pdf.cell(room, 5.4, dates, align='R')
                pdf.set_xy(pdf.l_margin, y)
            pdf.set_font('Helvetica', 'B', 10.5)
            pdf.multi_cell(width - room, 5.4, line, align='L', new_x=XPos.LMARGIN, new_y=YPos.NEXT)
        elif kind == 'org':
            pdf.set_font('Helvetica', 'I', 10.5)
            pdf.set_text_color(70)
            pdf.multi_cell(width, 4.9, line, align='L', new_x=XPos.LMARGIN, new_y=YPos.NEXT)
            pdf.set_text_color(0)
            pdf.ln(0.8)
        elif kind == 'name':
            pdf.set_font('Helvetica', 'B', 20)
            pdf.multi_cell(width, 9, line, align='C', new_x=XPos.LMARGIN, new_y=YPos.NEXT)
        elif kind == 'contact':
            pdf.set_font('Helvetica', '', 9.5)
            pdf.multi_cell(width, 5, line, align='C', new_x=XPos.LMARGIN, new_y=YPos.NEXT)
        elif kind == 'heading':
            if pdf.get_y() > pdf.page_break_trigger - 25:
                pdf.add_page()
            pdf.ln(4)
            pdf.set_font('Helvetica', 'B', 11)
            pdf.multi_cell(width, 6, line, align='L', new_x=XPos.LMARGIN, new_y=YPos.NEXT)
            pdf.set_draw_color(128)
            pdf.line(pdf.l_margin, pdf.get_y(), pdf.l_margin + width, pdf.get_y())
            pdf.ln(2)
        elif kind == 'bullet':
            pdf.set_font('Helvetica', '', 10.5)
            pdf.set_fill_color(40)
            pdf.circle(pdf.l_margin + 2, pdf.get_y() + 2.45, 1, style='F')
            pdf.set_x(pdf.l_margin + 5)
            pdf.multi_cell(width - 5, 4.9, line, align='L', new_x=XPos.LMARGIN, new_y=YPos.NEXT)
        elif kind == 'text':
            pdf.set_font('Helvetica', '', 10.5)
            pdf.multi_cell(width, 4.9, line, align='L', new_x=XPos.LMARGIN, new_y=YPos.NEXT)
        else:
            pdf.ln(1.5)
        previous = kind

    return bytes(pdf.output())
